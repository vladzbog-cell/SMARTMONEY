import asyncio
import mimetypes
import re
from typing import Any, cast

import aiohttp
import docx
import fitz  # type: ignore[import-not-found]
from bs4 import BeautifulSoup
from PIL import Image
from youtube_transcript_api import YouTubeTranscriptApi
from google.genai import types

from artemox_client import (
    ARTEMOX_MEDIA_MODEL,
    ARTEMOX_MEDIA_QUOTA_EXCEEDED_SENTINEL,
    get_artemox_client,
)


MEDIA_EXTRACTION_BRIEF = (
    "Ты работаешь как интеллектуальный редактор для сервиса создания презентаций уровня NotebookLM. "
    "Твоя задача — не просто распознать вход, а извлечь из него всё полезное для будущего storytelling.\n\n"
    "На выходе нужен максимально полезный структурированный контекст:\n"
    "1. Ключевая тема и цель материала.\n"
    "2. Основные факты, цифры, тезисы, аргументы и выводы.\n"
    "3. Скрытая логика: проблема, причины, последствия, решения, сравнения, сценарии, приоритеты.\n"
    "4. Если видишь таблицы, графики, схемы, интерфейсы, кластеры, карты, оси сравнения — подробно опиши их структуру.\n"
    "5. Если есть явная аудитория или тональность, назови их.\n"
    "6. Ничего не упускай, но убирай шум и повторы."
)


def _read_pdf_text(file_path: str) -> str:
    parts: list[str] = []
    with fitz.open(file_path) as document:
        for page in document:
            page_obj = cast(Any, page)
            text = str(page_obj.get_text("text")).strip()
            if text:
                parts.append(text)
    return "\n".join(parts)


def _extract_video_id(url: str) -> str | None:
    patterns = [
        r"(?:v=|/)([0-9A-Za-z_-]{11})(?:[?&/]|$)",
        r"youtu\.be/([0-9A-Za-z_-]{11})(?:[?&/]|$)",
    ]
    for pattern in patterns:
        match = re.search(pattern, url)
        if match:
            return str(match.group(1))
    return None


def _fetch_youtube_transcript(video_id: str) -> str:
    api = cast(Any, YouTubeTranscriptApi)
    transcript = cast(list[dict[str, Any]], api.get_transcript(video_id, languages=["ru", "en"]))
    return " ".join(str(item.get("text", "")).strip() for item in transcript if item.get("text"))


def _read_docx_text(file_path: str) -> str:
    try:
        document = docx.Document(file_path)
        full_text: list[str] = []
        for paragraph in document.paragraphs:
            try:
                text = paragraph.text.strip()
            except Exception:
                continue
            if text:
                full_text.append(text)

        # Also read table cells: many analytical DOCX files keep core facts in tables.
        for table in document.tables:
            for row in table.rows:
                row_cells = []
                for cell in row.cells:
                    value = re.sub(r"\s+", " ", cell.text or "").strip()
                    if value:
                        row_cells.append(value)
                if row_cells:
                    full_text.append(" | ".join(row_cells))
        return "\n".join(full_text)
    except Exception as e:
        print(f"Ошибка парсинга DOCX: {e}")
        return ""


def _extract_media_context(file_path: str, media_type: str) -> str:
    client = get_artemox_client()
    if client is None:
        return ""

    try:
        if media_type == "image":
            with Image.open(file_path) as img:
                response = client.models.generate_content(
                    model=ARTEMOX_MEDIA_MODEL,
                    contents=[
                        MEDIA_EXTRACTION_BRIEF
                        + "\n\nИзвлеки все данные, включая цифры из таблиц и зависимости на схемах. "
                        "Распознай абсолютно весь текст на изображении. "
                        "Если есть схемы, графики, таблицы, интерфейсы, карты, оси сравнения или заметки от руки — подробно опиши их структуру, смысл и возможную роль в презентации.",
                        img,
                    ],
                )
            return getattr(response, "text", "") or ""

        if media_type == "audio":
            mime_type = mimetypes.guess_type(file_path)[0] or "audio/mpeg"
            with open(file_path, "rb") as audio_file:
                audio_bytes = audio_file.read()
            audio_part = types.Part.from_bytes(data=audio_bytes, mime_type=mime_type)
            response = client.models.generate_content(
                model=ARTEMOX_MEDIA_MODEL,
                contents=[
                    MEDIA_EXTRACTION_BRIEF
                    + "\n\nСделай полную и точную расшифровку аудио, а затем выдели из неё полезный контекст для презентации: "
                    "главные идеи, аргументы, факты, цифры, формулировки выводов и возможные блоки будущих слайдов.",
                    audio_part,
                ],
            )
            return getattr(response, "text", "") or ""

        if media_type == "video":
            mime_type = mimetypes.guess_type(file_path)[0] or "video/mp4"
            with open(file_path, "rb") as video_file:
                video_bytes = video_file.read()
            video_part = types.Part.from_bytes(data=video_bytes, mime_type=mime_type)
            response = client.models.generate_content(
                model=ARTEMOX_MEDIA_MODEL,
                contents=[
                    MEDIA_EXTRACTION_BRIEF
                    + "\n\nПроанализируй видео и извлеки всё полезное для будущей презентации: "
                    "речь, текст в кадре, ключевые сцены, объекты, графики, интерфейсы, последовательность идей, спорные точки, выводы и скрытую структуру материала.",
                    video_part,
                ],
            )
            return getattr(response, "text", "") or ""

        return ""
    except Exception as e:
        print(f"🔥 Ошибка распознавания {media_type}: {e}")
        if "429" in str(e) or "quota" in str(e).lower() or "limit" in str(e).lower():
            return ARTEMOX_MEDIA_QUOTA_EXCEEDED_SENTINEL
        return ""


async def extract_text_from_pdf(file_path: str) -> str:
    try:
        return await asyncio.to_thread(_read_pdf_text, file_path)
    except Exception:
        return ""


async def extract_text_from_docx(file_path: str) -> str:
    try:
        return await asyncio.to_thread(_read_docx_text, file_path)
    except Exception:
        return ""


async def extract_text_from_url(url: str) -> str:
    try:
        timeout = aiohttp.ClientTimeout(total=30)
        async with aiohttp.ClientSession(timeout=timeout) as session:
            async with session.get(url) as response:
                if response.status != 200:
                    return ""
                html = await response.text()

        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()

        text = soup.get_text(separator=" ")
        return " ".join(text.split())
    except Exception:
        return ""


async def extract_text_from_youtube(url: str) -> str:
    try:
        video_id = _extract_video_id(url)
        if not video_id:
            return ""
        return await asyncio.to_thread(_fetch_youtube_transcript, video_id)
    except Exception:
        return ""


async def extract_context_from_media(file_path: str, media_type: str) -> str:
    try:
        return await asyncio.to_thread(_extract_media_context, file_path, media_type)
    except Exception:
        return ""
