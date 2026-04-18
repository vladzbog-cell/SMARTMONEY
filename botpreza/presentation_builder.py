import asyncio
import hashlib
import json
import os
import re

try:
    from PIL import Image, ImageStat  # type: ignore
    _PIL_AVAILABLE = True
except Exception:  # pragma: no cover
    _PIL_AVAILABLE = False

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from image_service import generate_slide_image

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = RGBColor(247, 246, 242)
SURFACE = RGBColor(255, 255, 252)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(30, 43, 67)
MUTED = RGBColor(97, 107, 126)
LINE = RGBColor(208, 214, 223)
ACCENT = RGBColor(50, 126, 151)
ACCENT_ALT = RGBColor(207, 128, 75)
ACCENT_SOFT = RGBColor(229, 236, 241)
OVERLAY_PANEL = RGBColor(30, 30, 30)
TITLE_ON_DARK = RGBColor(255, 255, 255)
BODY_ON_DARK = RGBColor(232, 236, 242)
CARD_BG = RGBColor(250, 252, 255)
CARD_TITLE = RGBColor(24, 36, 58)
CARD_BODY = RGBColor(72, 86, 112)

# Typography tokens — one source of truth for size + family hierarchy.
FONT_DISPLAY = "Georgia"   # editorial serif for hero/cover titles
FONT_BASE = "Arial"        # universal body / UI
FONT_MONO = "Consolas"     # rare: stats caption micro-text
TS_DISPLAY = 48            # huge hero title (cover panel)
TS_HEADLINE = 30           # main title in most layouts
TS_H1 = 22                 # section lead / quote body
TS_H2 = 16                 # subtitle / lead paragraph
TS_BODY = 12               # body copy
TS_SMALL = 10              # supporting / caption
TS_MICRO = 8               # footer / source badge
LINE_SPACING_BODY = 1.18   # breathing room for multi-line body copy

# Editorial canvas palette (NotebookLM-inspired)
CANVAS = RGBColor(247, 243, 233)            # warm cream background
HEADLINE_INK = RGBColor(18, 28, 52)         # deep navy serif headline
BODY_INK = RGBColor(44, 56, 78)             # primary body ink
SECONDARY_INK = RGBColor(110, 118, 132)     # kicker, meta, captions
HAIRLINE = RGBColor(204, 198, 184)          # thin rule under headline
COPPER = RGBColor(196, 108, 62)             # rare accent (numerals, highlight)
CARD_SURFACE = RGBColor(252, 248, 238)      # callout box fill
CARD_BORDER = RGBColor(214, 206, 190)       # callout box border
WATERMARK = RGBColor(160, 160, 150)         # bottom-right logo mark


def _clean_json_string(json_string: str) -> str:
    cleaned = (json_string or "").strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        if cleaned.lower().startswith("json"):
            cleaned = cleaned[4:].strip()
    start_index = -1
    for marker in ("[", "{"):
        index = cleaned.find(marker)
        if index != -1 and (start_index == -1 or index < start_index):
            start_index = index
    if start_index > 0:
        cleaned = cleaned[start_index:]
    end_index = max(cleaned.rfind("]"), cleaned.rfind("}"))
    if end_index != -1:
        cleaned = cleaned[: end_index + 1]
    return cleaned.strip()


def _normalize_slides(payload) -> list[dict]:
    if isinstance(payload, list):
        return [s for s in payload if isinstance(s, dict)]
    if isinstance(payload, dict):
        slides = payload.get("slides")
        if isinstance(slides, list):
            return [s for s in slides if isinstance(s, dict)]
        if "image_prompt" in payload:
            return [payload]
    raise ValueError("Unsupported presentation payload format")


def _extract_title(slide_data: dict, fallback_number: int) -> str:
    content_block = slide_data.get("content")
    if isinstance(content_block, dict):
        for key in ("title", "headline", "heading", "name", "topic"):
            value = str(content_block.get(key) or "").strip()
            if value and value.lower() not in ("null", "none"):
                normalized = re.sub(r"\s+", " ", value).strip(" .;:-")
                if normalized:
                    return normalized[:84]

    for key in ("title", "headline", "heading", "slide_title", "name", "topic"):
        value = str(slide_data.get(key) or "").strip()
        if value and value.lower() not in ("null", "none"):
            normalized = re.sub(r"\s+", " ", value).strip(" .;:-")
            if len(normalized) > 84:
                cut = normalized[:84]
                pivot = max(cut.rfind("."), cut.rfind(":"), cut.rfind(","), cut.rfind(" "))
                if pivot > 52:
                    cut = cut[:pivot]
                normalized = cut.strip(" .;:-")
            return normalized
    return f"Слайд {fallback_number}"


def _normalize_bullet_lines(value) -> list[str]:
    if isinstance(value, list):
        result = []
        for item in value:
            if isinstance(item, dict):
                nested = ""
                for key in ("text", "value", "label", "title", "body", "content", "point"):
                    candidate = item.get(key)
                    if isinstance(candidate, str) and candidate.strip():
                        nested = candidate.strip()
                        break
                if not nested:
                    continue
                item = nested
            elif isinstance(item, (list, tuple)):
                item = " ".join(str(x) for x in item if isinstance(x, (str, int, float)))
            text = str(item or "").strip()
            if text and text.lower() not in ("null", "none") and not text.startswith("{"):
                cleaned = re.sub(r"^\d+\.\s*", "", text.lstrip("-• ").strip())
                cleaned = re.sub(r"\s+", " ", cleaned).strip(" .;:-")
                if cleaned:
                    result.append(cleaned)
        return result

    text = str(value or "").strip()
    if not text or text.lower() in ("null", "none"):
        return []

    result = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if line:
            cleaned = re.sub(r"^\d+\.\s*", "", line.lstrip("-• ").strip())
            cleaned = re.sub(r"\s+", " ", cleaned).strip(" .;:-")
            result.append(cleaned)
    return result


def _dedupe_and_shorten(lines: list[str], *, max_len: int = 128, max_items: int = 4) -> list[str]:
    result: list[str] = []
    seen: list[str] = []
    for line in lines:
        text = re.sub(r"\(([^)]{1,80})\)\s*\(\1\)", r"(\1)", line)
        text = re.sub(r"\s+", " ", text).strip(" .;:-")
        if not text:
            continue
        if len(text) > max_len:
            cut = text[:max_len]
            pivot = max(cut.rfind("."), cut.rfind(":"), cut.rfind(","), cut.rfind(" "))
            if pivot > int(max_len * 0.6):
                cut = cut[:pivot]
            text = cut.strip(" .;:-")
        key = re.sub(r"[^a-zа-яё0-9 ]+", " ", text.lower())
        key = re.sub(r"\s+", " ", key).strip()
        if len(key) < 8:
            continue
        if key in seen:
            continue
        if any((key in prev or prev in key) and min(len(key), len(prev)) > 20 for prev in seen):
            continue
        result.append(text)
        seen.append(key)
        if len(result) >= max_items:
            break
    return result


def _extract_bullets(slide_data: dict) -> list[str]:
    content_block = slide_data.get("content")
    if isinstance(content_block, dict):
        for key in ("bullets", "theses", "points", "key_points", "highlights", "items", "text"):
            lines = _dedupe_and_shorten(_normalize_bullet_lines(content_block.get(key)))
            if lines:
                return lines[:6]

    for key in ("bullets", "theses", "points", "key_points", "highlights", "content", "body", "text"):
        lines = _dedupe_and_shorten(_normalize_bullet_lines(slide_data.get(key)))
        if lines:
            return lines[:6]

    subtitle = str(slide_data.get("subtitle") or "").strip()
    if subtitle and subtitle.lower() not in ("null", "none"):
        return [subtitle]

    image_prompt = str(slide_data.get("image_prompt") or "").strip()
    if image_prompt:
        fallback = image_prompt[:180] + "..." if len(image_prompt) > 180 else image_prompt
        return _dedupe_and_shorten([fallback], max_items=1)
    return []


def _kicker(layout: str) -> str:
    mapping = {
        "hero": "Стратегический обзор",
        "cards": "Ключевые темы",
        "process": "Последовательность действий",
        "chart_focus": "Динамика и метрики",
        "comparison": "Сравнение сценариев",
        "matrix": "Матрица приоритетов",
        "radial": "Системная архитектура",
        "cover": "Обложка",
        "agenda": "Повестка",
        "context": "Контекст",
        "key_findings": "Ключевые находки",
        "stat_highlight": "Главная цифра",
        "quote_highlight": "Цитата из источника",
        "takeaways": "Выводы и действия",
        "hero_concept": "Стратегический обзор",
        "split_infographic": "Ключевые темы",
        "timeline": "Последовательность действий",
        "data_matrix": "Матрица приоритетов",
    }
    return mapping.get(layout, "Аналитический слайд")


def _extract_subtitle(slide_data: dict) -> str:
    content_block = slide_data.get("content") if isinstance(slide_data.get("content"), dict) else {}
    for source in (content_block, slide_data):
        for key in ("subtitle", "lead", "dek", "kicker_text"):
            value = str(source.get(key) or "").strip()
            if value and value.lower() not in ("null", "none"):
                return re.sub(r"\s+", " ", value).strip(" .;:-")[:200]
    return ""


def _extract_stats(slide_data: dict) -> list[dict]:
    content_block = slide_data.get("content") if isinstance(slide_data.get("content"), dict) else {}
    raw = None
    for source in (content_block, slide_data):
        for key in ("stats", "numbers", "metrics"):
            candidate = source.get(key)
            if isinstance(candidate, list) and candidate:
                raw = candidate
                break
        if raw is not None:
            break
    if not raw:
        return []
    stats: list[dict] = []
    for item in raw:
        if isinstance(item, dict):
            value = re.sub(r"\s+", " ", str(item.get("value") or item.get("number") or "")).strip()[:24]
            label = re.sub(r"\s+", " ", str(item.get("label") or item.get("caption") or item.get("description") or "")).strip()[:90]
        else:
            text = re.sub(r"\s+", " ", str(item or "")).strip()
            match = re.match(r"(\S+?)\s*[-—:]\s*(.+)", text)
            if match:
                value, label = match.group(1)[:24], match.group(2)[:90]
            else:
                value, label = "", text[:90]
        if value or label:
            stats.append({"value": value, "label": label})
        if len(stats) >= 4:
            break
    return stats


def _extract_quotes(slide_data: dict) -> list[dict]:
    content_block = slide_data.get("content") if isinstance(slide_data.get("content"), dict) else {}
    raw = None
    for source in (content_block, slide_data):
        for key in ("quotes", "citations"):
            candidate = source.get(key)
            if isinstance(candidate, list) and candidate:
                raw = candidate
                break
        if raw is not None:
            break
    if not raw:
        return []
    quotes: list[dict] = []
    for item in raw:
        if isinstance(item, dict):
            text = re.sub(r"\s+", " ", str(item.get("text") or item.get("quote") or "")).strip()[:240]
            author = re.sub(r"\s+", " ", str(item.get("author") or item.get("source") or "")).strip()[:90]
        else:
            text = re.sub(r"\s+", " ", str(item or "")).strip()[:240]
            author = ""
        if text:
            quotes.append({"text": text, "author": author})
        if len(quotes) >= 3:
            break
    return quotes


def _extract_source_hint(slide_data: dict) -> str:
    content_block = slide_data.get("content") if isinstance(slide_data.get("content"), dict) else {}
    for source in (content_block, slide_data):
        for key in ("source_hint", "source", "provenance"):
            value = str(source.get(key) or "").strip()
            if value and value.lower() not in ("null", "none"):
                return re.sub(r"\s+", " ", value).strip(" .;:-")[:220]
    return ""


def _extract_speaker_notes(slide_data: dict) -> str:
    content_block = slide_data.get("content") if isinstance(slide_data.get("content"), dict) else {}
    for source in (slide_data, content_block):
        value = source.get("speaker_notes") or source.get("notes")
        if isinstance(value, list):
            joined = " ".join(str(item or "").strip() for item in value if str(item or "").strip())
        else:
            joined = str(value or "").strip()
        joined = re.sub(r"\s+", " ", joined).strip()
        if joined and joined.lower() not in ("null", "none"):
            return joined[:1800]
    return ""


def _compose_notes_text(slide_data: dict, title: str, layout: str, bullets: list[str]) -> str:
    parts: list[str] = []
    notes = _extract_speaker_notes(slide_data)
    if notes:
        parts.append(notes)
    else:
        fallback_lead = {
            "cover": "Открываем презентацию и обозначаем тему.",
            "agenda": "Озвучиваем повестку: какие блоки рассмотрим.",
            "context": "Даём контекст и рамку темы.",
            "key_findings": "Перечисляем ключевые находки.",
            "stat_highlight": "Акцентируем внимание на главной цифре.",
            "quote_highlight": "Усиляем идею цитатой из источника.",
            "split_infographic": "Разбираем блок карточек по очереди.",
            "timeline": "Проходим по этапам последовательно.",
            "data_matrix": "Сравниваем план и реальность.",
            "comparison": "Сравниваем два сценария.",
            "takeaways": "Сводим выводы и следующие шаги.",
            "hero_concept": "Формулируем ключевой концепт слайда.",
        }.get(layout, "Комментируем слайд своими словами.")
        bullet_hint = "; ".join(item for item in bullets[:3]) if bullets else ""
        parts.append(f"{fallback_lead} Тема: «{title}».")
        if bullet_hint:
            parts.append(f"Тезисы: {bullet_hint}.")

    source_hint = _extract_source_hint(slide_data)
    if source_hint:
        parts.append(f"Опора: {source_hint}")

    return "\n\n".join(parts)[:1800]


def _attach_speaker_notes(slide, notes_text: str) -> None:
    if not notes_text:
        return
    try:
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = notes_text
    except Exception as exc:  # pragma: no cover - defensive
        print(f"⚠️ Failed to attach speaker notes: {exc}")


def _layout_type(slide_data: dict, slide_number: int) -> str:
    raw = str(slide_data.get("layout_type") or "").strip().lower()
    aliases = {
        "split_infographic": "split_infographic",
        "hero_concept": "hero_concept",
        "data_matrix": "data_matrix",
        "timeline": "timeline",
        "hero": "hero_concept",
        "cards": "split_infographic",
        "process": "timeline",
        "chart_focus": "split_infographic",
        "matrix": "data_matrix",
        "radial": "split_infographic",
        "cover": "cover",
        "title": "cover",
        "opening": "cover",
        "agenda": "agenda",
        "toc": "agenda",
        "outline": "agenda",
        "context": "context",
        "intro": "context",
        "key_findings": "key_findings",
        "findings": "key_findings",
        "insights": "key_findings",
        "stat_highlight": "stat_highlight",
        "stat": "stat_highlight",
        "number": "stat_highlight",
        "metric": "stat_highlight",
        "quote_highlight": "quote_highlight",
        "quote": "quote_highlight",
        "pull_quote": "quote_highlight",
        "takeaways": "takeaways",
        "summary": "takeaways",
        "conclusion": "takeaways",
        "closing": "takeaways",
        "comparison": "comparison",
        "versus": "comparison",
    }
    if raw in aliases:
        return aliases[raw]
    cycle = ["cover", "agenda", "context", "key_findings", "split_infographic", "timeline", "data_matrix", "takeaways"]
    return cycle[(slide_number - 1) % len(cycle)]


def _hash_values(seed: str, count: int, min_val: int, max_val: int) -> list[int]:
    digest = hashlib.sha256(seed.encode("utf-8")).digest()
    span = max_val - min_val
    values = []
    for idx in range(count):
        values.append(min_val + digest[idx] % max(span, 1))
    return values


def _set_slide_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = CANVAS


def _add_grid(slide, left, top, width, height, step_x=Inches(0.42), step_y=Inches(0.34), opacity=0.55):
    line_color = RGBColor(228, 232, 238)
    x = left
    while x <= left + width:
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, top, x, top + height)
        line.line.color.rgb = line_color
        line.line.transparency = opacity
        line.line.width = Pt(0.5)
        x += step_x
    y = top
    while y <= top + height:
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, left, y, left + width, y)
        line.line.color.rgb = line_color
        line.line.transparency = opacity
        line.line.width = Pt(0.5)
        y += step_y


def _add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    size,
    color,
    *,
    bold=False,
    align=PP_ALIGN.LEFT,
    font_name: str = FONT_BASE,
    line_spacing: float | None = None,
    italic: bool = False,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    if line_spacing is not None:
        paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def _add_panel(slide, left, top, width, height, *, radius=True, fill_color=SURFACE, transparency=0.1):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.fill.transparency = transparency
    shape.line.color.rgb = LINE
    shape.line.width = Pt(0.8)
    return shape


def _add_full_bleed_background(slide, image_path: str, *, veil_color=WHITE, veil_transparency=0.18) -> None:
    slide.shapes.add_picture(image_path, 0, 0, width=SLIDE_W, height=SLIDE_H)
    veil = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    veil.fill.solid()
    veil.fill.fore_color.rgb = veil_color
    veil.fill.transparency = veil_transparency
    veil.line.fill.background()


def _apply_full_bleed_base(slide, image_path: str, *, veil_transparency=0.26) -> None:
    _add_full_bleed_background(slide, image_path, veil_color=WHITE, veil_transparency=veil_transparency)


def _add_accent_bar(slide, left, top, width, color):
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar


def _add_accent_image(slide, image_path: str, left, top, width, height) -> None:
    slide.shapes.add_picture(image_path, left, top, width=width, height=height)
    frame = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    frame.fill.background()
    frame.line.color.rgb = LINE
    frame.line.width = Pt(0.6)


def _add_editorial_image_accent(slide, image_path: str, left, top, width, height, *, veil_transparency=0.78) -> None:
    slide.shapes.add_picture(image_path, left, top, width=width, height=height)
    frame = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    frame.fill.background()
    frame.line.color.rgb = LINE
    frame.line.width = Pt(0.7)


def _add_hairline_rule(slide, x1, y1, x2, y2, *, color=LINE, width=0.8, transparency=0.08):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.transparency = transparency
    return line


def _split_label(text: str) -> tuple[str, str]:
    text = (text or "").strip()
    if ":" in text:
        head, tail = text.split(":", 1)
        if len(head) < 40:
            return head.strip(), tail.strip()
    words = text.split()
    if len(words) > 5:
        return " ".join(words[:3]), " ".join(words[3:])
    return text, ""


def _draw_title_band(slide, title: str, kicker: str | None = None) -> None:
    if kicker:
        _add_textbox(slide, Inches(0.72), Inches(0.46), Inches(2.5), Inches(0.18), kicker.upper(), 8, ACCENT_ALT, bold=True)
    _add_hairline_rule(slide, Inches(0.72), Inches(0.66), Inches(12.45), Inches(0.66), transparency=0.0)
    _add_textbox(slide, Inches(0.72), Inches(0.82), Inches(7.9), Inches(0.58), title, 24, INK, bold=True)


def _render_hero(slide, image_path: str, title: str, bullets: list[str]) -> None:
    _apply_full_bleed_base(slide, image_path, veil_transparency=0.24)
    _add_hairline_rule(slide, Inches(0.55), Inches(0.55), Inches(12.8), Inches(0.55), transparency=0.0)
    _add_panel(slide, Inches(0.58), Inches(0.82), Inches(5.5), Inches(5.95), fill_color=WHITE, transparency=0.14)
    _add_hairline_rule(slide, Inches(6.55), Inches(1.02), Inches(6.55), Inches(6.62), transparency=0.0)
    _add_hairline_rule(slide, Inches(6.55), Inches(6.62), Inches(12.5), Inches(6.62), transparency=0.0)
    _add_textbox(slide, Inches(0.72), Inches(0.86), Inches(4.9), Inches(0.22), _kicker("hero").upper(), 8, ACCENT_ALT, bold=True)
    _add_textbox(slide, Inches(0.72), Inches(1.18), Inches(5.15), Inches(2.75), title, 31, INK, bold=True)
    subtitle = bullets[0] if bullets else ""
    if subtitle:
        _add_textbox(slide, Inches(0.72), Inches(3.95), Inches(4.95), Inches(0.88), subtitle, 17, INK)
    for idx, bullet in enumerate(bullets[1:3], start=0):
        top = Inches(5.0 + idx * 0.66)
        _add_hairline_rule(slide, Inches(0.72), top + Inches(0.18), Inches(0.98), top + Inches(0.18), color=ACCENT if idx == 0 else ACCENT_ALT, width=2.2, transparency=0.0)
        _add_textbox(slide, Inches(1.08), top, Inches(4.6), Inches(0.34), bullet, 11, MUTED)
    # No system footer: keeps title slide cleaner and closer to editorial reference style.


def _render_cards(slide, image_path: str, title: str, bullets: list[str]) -> None:
    _apply_full_bleed_base(slide, image_path, veil_transparency=0.27)
    _draw_title_band(slide, title, _kicker("cards"))
    positions = [
        (Inches(0.72), Inches(2.0), Inches(5.55), Inches(1.78)),
        (Inches(6.0), Inches(2.0), Inches(6.0), Inches(1.78)),
        (Inches(0.72), Inches(4.08), Inches(5.55), Inches(1.78)),
        (Inches(6.0), Inches(4.08), Inches(6.0), Inches(1.78)),
    ]
    for idx, bullet in enumerate((bullets + [""] * 4)[:4]):
        heading, body = _split_label(bullet or "Insight")
        left, top, width, height = positions[idx]
        fill = ACCENT_SOFT if idx in (0, 3) else SURFACE
        _add_panel(slide, left, top, width, height, fill_color=fill, transparency=0.03)
        _add_hairline_rule(slide, left, top + Inches(0.14), left + Inches(0.86), top + Inches(0.14), color=ACCENT if idx % 2 == 0 else ACCENT_ALT, width=2.2, transparency=0.0)
        _add_textbox(slide, left + Inches(0.2), top + Inches(0.28), width - Inches(0.42), Inches(0.28), heading, 13, INK, bold=True)
        _add_textbox(slide, left + Inches(0.2), top + Inches(0.66), width - Inches(0.42), Inches(0.68), body or heading, 11, MUTED)


def _render_process(slide, image_path: str, title: str, bullets: list[str]) -> None:
    _apply_full_bleed_base(slide, image_path, veil_transparency=0.27)
    _draw_title_band(slide, title, _kicker("process"))
    fallback_steps = ["Диагностика", "Приоритизация", "Внедрение", "Контроль результата"]
    fallback_step_bodies = [
        "Собрать факты, ограничения и контекст задачи.",
        "Определить приоритеты по влиянию и сложности.",
        "Запустить меры и закрепить зоны ответственности.",
        "Контролировать KPI и корректировать план.",
    ]
    steps = (bullets + fallback_steps)[:4]
    start_x = Inches(0.65)
    y = Inches(3.45)
    step_w = Inches(2.8)
    _add_hairline_rule(slide, Inches(1.05), Inches(4.32), Inches(12.05), Inches(4.32), transparency=0.0)
    for idx, bullet in enumerate(steps):
        left = start_x + idx * Inches(3.05)
        heading, body = _split_label(bullet)
        if not body or body.lower() == heading.lower():
            body = fallback_step_bodies[idx] if idx < len(fallback_step_bodies) else "Ключевые действия этапа"
        _add_panel(slide, left, y, step_w, Inches(1.45), fill_color=WHITE, transparency=0.02)
        chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left + Inches(0.14), y - Inches(0.32), Inches(0.5), Inches(0.5))
        chip.fill.solid()
        chip.fill.fore_color.rgb = ACCENT if idx % 2 == 0 else ACCENT_ALT
        chip.line.fill.background()
        _add_textbox(slide, left + Inches(0.22), y + Inches(0.16), Inches(2.3), Inches(0.26), f"{idx + 1}. {heading}", 13, INK, bold=True)
        _add_textbox(slide, left + Inches(0.22), y + Inches(0.5), Inches(2.3), Inches(0.5), body or heading, 10, MUTED)
        if idx < len(steps) - 1:
            _add_hairline_rule(slide, left + step_w, y + Inches(0.74), left + Inches(3.0), y + Inches(0.74), transparency=0.0)


def _render_chart_focus(slide, image_path: str, title: str, bullets: list[str]) -> None:
    _apply_full_bleed_base(slide, image_path, veil_transparency=0.28)
    _draw_title_band(slide, title, _kicker("chart_focus"))
    _add_panel(slide, Inches(0.6), Inches(1.75), Inches(7.55), Inches(4.95))
    _add_grid(slide, Inches(1.0), Inches(2.2), Inches(6.7), Inches(3.95), step_x=Inches(0.85), step_y=Inches(0.72), opacity=0.35)
    values = _hash_values(title + "|" + " ".join(bullets), 6, 10, 90)
    points = []
    for idx, value in enumerate(values):
        x = Inches(1.12 + idx * 1.08)
        y = Inches(5.8 - (value / 100) * 3.0)
        points.append((x, y))
    for idx in range(len(points) - 1):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, points[idx][0], points[idx][1], points[idx + 1][0], points[idx + 1][1])
        line.line.color.rgb = ACCENT if idx < len(points) - 2 else ACCENT_ALT
        line.line.width = Pt(2.0)
    for x, y in points:
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x - Inches(0.06), y - Inches(0.06), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT
        dot.line.fill.background()
    for idx, bullet in enumerate((bullets + [""] * 2)[:2]):
        heading, body = _split_label(bullet or "Finding")
        top = Inches(4.1 + idx * 1.2)
        _add_panel(slide, Inches(8.4), top, Inches(4.3), Inches(1.05))
        _add_textbox(slide, Inches(8.62), top + Inches(0.16), Inches(3.8), Inches(0.25), heading, 12, INK, bold=True)
        _add_textbox(slide, Inches(8.62), top + Inches(0.48), Inches(3.8), Inches(0.35), body or heading, 10, MUTED)


def _render_comparison(slide, image_path: str, title: str, bullets: list[str]) -> None:
    _apply_full_bleed_base(slide, image_path, veil_transparency=0.27)
    _draw_title_band(slide, title, _kicker("comparison"))
    _add_panel(slide, Inches(0.6), Inches(1.95), Inches(5.85), Inches(4.7))
    _add_panel(slide, Inches(6.85), Inches(1.95), Inches(5.85), Inches(4.7))
    left_title, left_body = _split_label(bullets[0] if bullets else "Левая позиция")
    right_title, right_body = _split_label(bullets[1] if len(bullets) > 1 else "Правая позиция")
    _add_accent_bar(slide, Inches(0.6), Inches(1.95), Inches(0.8), ACCENT)
    _add_accent_bar(slide, Inches(6.85), Inches(1.95), Inches(0.8), ACCENT_ALT)
    _add_textbox(slide, Inches(0.85), Inches(2.25), Inches(5.1), Inches(0.35), left_title, 14, INK, bold=True)
    _add_textbox(slide, Inches(0.85), Inches(2.72), Inches(5.1), Inches(2.0), left_body or left_title, 12, MUTED)
    _add_textbox(slide, Inches(7.1), Inches(2.25), Inches(5.1), Inches(0.35), right_title, 14, INK, bold=True)
    _add_textbox(slide, Inches(7.1), Inches(2.72), Inches(5.1), Inches(2.0), right_body or right_title, 12, MUTED)
    if len(bullets) > 2:
        _add_panel(slide, Inches(3.95), Inches(5.9), Inches(5.45), Inches(0.62), fill_color=ACCENT_SOFT)
        _add_textbox(slide, Inches(4.15), Inches(6.08), Inches(5.0), Inches(0.2), bullets[2], 11, INK, align=PP_ALIGN.CENTER)


def _render_matrix(slide, image_path: str, title: str, bullets: list[str]) -> None:
    _apply_full_bleed_base(slide, image_path, veil_transparency=0.28)
    _draw_title_band(slide, title, _kicker("matrix"))
    _add_panel(slide, Inches(1.0), Inches(1.9), Inches(9.9), Inches(4.9))
    cx = Inches(5.95)
    cy = Inches(4.35)
    vertical = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cx, Inches(2.18), cx, Inches(6.45))
    vertical.line.color.rgb = LINE
    horizontal = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.28), cy, Inches(10.62), cy)
    horizontal.line.color.rgb = LINE
    axis_x = _add_textbox(slide, Inches(4.4), Inches(6.5), Inches(3.2), Inches(0.2), "Стратегическая привлекательность", 10, MUTED, align=PP_ALIGN.CENTER)
    axis_y = _add_textbox(slide, Inches(0.45), Inches(3.7), Inches(0.4), Inches(1.0), "Операционная сложность", 10, MUTED, align=PP_ALIGN.CENTER)
    axis_y.rotation = 270
    positions = [
        (Inches(1.35), Inches(2.3)),
        (Inches(6.18), Inches(2.3)),
        (Inches(1.35), Inches(4.78)),
        (Inches(6.18), Inches(4.78)),
    ]
    matrix_defaults = ["Быстрые победы", "Стратегические ставки", "Операционные улучшения", "Зона контроля"]
    for idx, bullet in enumerate((bullets + matrix_defaults)[:4]):
        heading, body = _split_label(bullet or matrix_defaults[idx])
        if not body or body.lower() == heading.lower():
            body = "Приоритет и критерии реализации"
        left, top = positions[idx]
        tint = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left - Inches(0.08), top - Inches(0.08), Inches(3.98), Inches(1.02))
        tint.fill.solid()
        tint.fill.fore_color.rgb = ACCENT_SOFT if idx in (0, 3) else WHITE
        tint.fill.transparency = 0.08
        tint.line.fill.background()
        _add_textbox(slide, left, top, Inches(3.9), Inches(0.3), heading, 13, INK, bold=True)
        _add_textbox(slide, left, top + Inches(0.32), Inches(3.9), Inches(0.62), body or heading, 10, MUTED)


def _render_radial(slide, image_path: str, title: str, bullets: list[str]) -> None:
    _apply_full_bleed_base(slide, image_path, veil_transparency=0.27)
    _draw_title_band(slide, title, _kicker("radial"))
    center_x = Inches(8.25)
    center_y = Inches(4.35)
    for radius, color in ((1.45, LINE), (1.05, ACCENT), (0.62, ACCENT_ALT)):
        ring = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, center_x - Inches(radius), center_y - Inches(radius), Inches(radius * 2), Inches(radius * 2))
        ring.fill.background()
        ring.line.color.rgb = color
        ring.line.width = Pt(1.2)
    core = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, center_x - Inches(0.44), center_y - Inches(0.44), Inches(0.88), Inches(0.88))
    core.fill.solid()
    core.fill.fore_color.rgb = ACCENT
    core.line.fill.background()
    _add_textbox(slide, center_x - Inches(0.6), center_y - Inches(0.1), Inches(1.2), Inches(0.2), "Ядро", 12, WHITE, bold=True, align=PP_ALIGN.CENTER)
    positions = [
        (Inches(4.15), Inches(1.95)),
        (Inches(9.6), Inches(1.95)),
        (Inches(4.15), Inches(4.95)),
        (Inches(9.6), Inches(4.95)),
    ]
    for idx, bullet in enumerate((bullets + [""] * 4)[:4]):
        heading, body = _split_label(bullet or "Node")
        left, top = positions[idx]
        _add_panel(slide, left, top, Inches(3.1), Inches(1.35))
        _add_textbox(slide, left + Inches(0.18), top + Inches(0.16), Inches(2.7), Inches(0.24), heading, 12, INK, bold=True)
        _add_textbox(slide, left + Inches(0.18), top + Inches(0.46), Inches(2.7), Inches(0.45), body or heading, 10, MUTED)


def _add_gradient_fallback_background(slide, seed: str) -> None:
    values = _hash_values(seed or "fallback", 9, 0, 255)
    c1 = RGBColor(18 + values[0] % 42, 32 + values[1] % 50, 60 + values[2] % 70)
    c2 = RGBColor(28 + values[3] % 54, 70 + values[4] % 72, 112 + values[5] % 72)
    c3 = RGBColor(116 + values[6] % 60, 96 + values[7] % 70, 72 + values[8] % 70)

    base = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    base.fill.solid()
    base.fill.fore_color.rgb = c1
    base.line.fill.background()

    glow_right = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        int(SLIDE_W * 0.47),
        int(SLIDE_H * -0.08),
        int(SLIDE_W * 0.7),
        int(SLIDE_H * 1.08),
    )
    glow_right.fill.solid()
    glow_right.fill.fore_color.rgb = c2
    glow_right.fill.transparency = 0.42
    glow_right.line.fill.background()

    glow_bottom = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        int(SLIDE_W * 0.45),
        int(SLIDE_H * 0.42),
        int(SLIDE_W * 0.65),
        int(SLIDE_H * 0.7),
    )
    glow_bottom.fill.solid()
    glow_bottom.fill.fore_color.rgb = c3
    glow_bottom.fill.transparency = 0.52
    glow_bottom.line.fill.background()


def _add_background_layer(slide, image_path: str | None, seed: str) -> None:
    if image_path and os.path.exists(image_path):
        slide.shapes.add_picture(image_path, 0, 0, width=SLIDE_W, height=SLIDE_H)
        return
    _add_gradient_fallback_background(slide, seed)


def _add_glass_panel(slide, *, width_ratio: float = 0.36, transparency: float = 0.32) -> int:
    panel_width = int(SLIDE_W * width_ratio)
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, panel_width, SLIDE_H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = OVERLAY_PANEL
    panel.fill.transparency = transparency
    panel.line.fill.background()

    divider = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        panel_width - Inches(0.02),
        0,
        Inches(0.02),
        SLIDE_H,
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor(255, 255, 255)
    divider.fill.transparency = 0.78
    divider.line.fill.background()
    return panel_width


def _ui_trim(text: str, limit: int, *, ellipsis: bool = True) -> str:
    normalized = re.sub(r"\s+", " ", (text or "").strip()).strip(" .;:-")
    if len(normalized) <= limit:
        return normalized
    cut = normalized[:limit]
    pivot = max(cut.rfind("."), cut.rfind(":"), cut.rfind(","), cut.rfind(" "))
    if pivot > int(limit * 0.55):
        cut = cut[:pivot]
    cut = cut.strip(" .;:-")
    if ellipsis and cut and not cut.endswith(("…", "...", ".")):
        cut = f"{cut}…"
    return cut


def _overlay_bullets_for_left(layout: str, bullets: list[str]) -> list[str]:
    base = _dedupe_and_shorten(bullets, max_len=108, max_items=6)
    if not base:
        return ["Ключевые тезисы слайда", "Контекст, выводы и действия"]

    if layout == "timeline":
        mapped = [_split_label(_ui_trim(item, 86))[0] for item in base]
        return _dedupe_and_shorten(mapped, max_len=70, max_items=3)

    if layout == "data_matrix":
        summary = []
        for item in base[:4]:
            plan, reality = _split_matrix_row(item)
            summary.append(f"{_ui_trim(plan, 42)} vs {_ui_trim(reality, 42)}")
        return _dedupe_and_shorten(summary, max_len=88, max_items=3)

    return _dedupe_and_shorten(base, max_len=96, max_items=3)


def _render_left_text_overlay(slide, title: str, bullets: list[str], panel_width: int, layout: str) -> None:
    content_left = Inches(0.62)
    content_width = panel_width - Inches(1.0)

    title_text = _ui_trim(title, 88, ellipsis=False)
    title_size = 32 if len(title_text) > 58 else 36

    _add_textbox(
        slide,
        content_left,
        Inches(0.72),
        content_width,
        Inches(2.3),
        title_text,
        title_size,
        TITLE_ON_DARK,
        bold=True,
        font_name=FONT_DISPLAY,
        line_spacing=1.08,
    )

    normalized = _overlay_bullets_for_left(layout, bullets)
    if not normalized:
        normalized = ["Ключевые тезисы слайда", "Контекст, выводы и действия"]

    for idx, bullet in enumerate(normalized):
        _add_textbox(
            slide,
            content_left,
            Inches(3.0 + idx * 0.95),
            content_width,
            Inches(0.82),
            f"• {_ui_trim(bullet, 96)}",
            TS_H2,
            BODY_ON_DARK,
            line_spacing=LINE_SPACING_BODY,
        )


def _draw_split_infographic_right(slide, right_left, bullets: list[str]) -> None:
    card_w = Inches(2.78)
    card_h = Inches(1.62)
    positions = [
        (right_left + Inches(0.25), Inches(1.05)),
        (right_left + Inches(3.25), Inches(1.05)),
        (right_left + Inches(0.25), Inches(3.05)),
        (right_left + Inches(3.25), Inches(3.05)),
    ]
    items = (_dedupe_and_shorten(bullets, max_len=92, max_items=4) + ["Insight", "Фактор", "Риск", "Решение"])[:4]
    for idx, text in enumerate(items):
        left, top = positions[idx]
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.fill.transparency = 0.16
        card.line.color.rgb = RGBColor(255, 255, 255)
        card.line.transparency = 0.38
        card.line.width = Pt(1.0)
        accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left + Inches(0.14), top + Inches(0.14), Inches(0.72), Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = ACCENT if idx % 2 == 0 else ACCENT_ALT
        accent.line.fill.background()
        heading, body = _split_label(text)
        _add_textbox(slide, left + Inches(0.2), top + Inches(0.3), card_w - Inches(0.35), Inches(0.38), _ui_trim(heading, 48, ellipsis=False), 13, CARD_TITLE, bold=True)
        _add_textbox(slide, left + Inches(0.2), top + Inches(0.68), card_w - Inches(0.35), Inches(0.62), _ui_trim(body or heading, 72), 10, CARD_BODY)


def _draw_timeline_right(slide, right_left, bullets: list[str]) -> None:
    line_x = right_left + Inches(0.85)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, line_x, Inches(1.05), line_x, Inches(6.38))
    line.line.color.rgb = ACCENT
    line.line.transparency = 0.2
    line.line.width = Pt(2.0)

    steps = (_dedupe_and_shorten(bullets, max_len=92, max_items=4) + ["Диагностика", "Приоритизация", "Запуск", "Контроль"])[:4]
    for idx, step in enumerate(steps):
        heading, body = _split_label(step)
        y = Inches(1.38 + idx * 1.3)
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, line_x - Inches(0.1), y - Inches(0.1), Inches(0.2), Inches(0.2))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT if idx % 2 == 0 else ACCENT_ALT
        dot.line.fill.background()

        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            right_left + Inches(1.15),
            y - Inches(0.3),
            Inches(4.35),
            Inches(0.9),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.fill.transparency = 0.16
        card.line.color.rgb = WHITE
        card.line.transparency = 0.34
        card.line.width = Pt(0.8)
        _add_textbox(slide, right_left + Inches(1.33), y - Inches(0.16), Inches(3.95), Inches(0.34), f"{idx + 1}. {_ui_trim(heading, 52, ellipsis=False)}", 12, CARD_TITLE, bold=True)
        _add_textbox(slide, right_left + Inches(1.33), y + Inches(0.16), Inches(3.95), Inches(0.3), _ui_trim(body or heading, 64), 9, CARD_BODY)


def _split_matrix_row(text: str) -> tuple[str, str]:
    parts = re.split(r"\s+(?:vs|VS|против)\s+|[|]|→|->|—", text or "", maxsplit=1)
    if len(parts) == 2:
        left = parts[0].strip(" .;:-")
        right = parts[1].strip(" .;:-")
        if left and right:
            return left, right
    heading, body = _split_label(text)
    if body and body.lower() != heading.lower():
        return heading, body
    return heading or "План", "Реальность"


def _draw_data_matrix_right(slide, right_left, bullets: list[str]) -> None:
    grid_left = right_left + Inches(0.25)
    grid_top = Inches(1.24)
    grid_w = Inches(5.85)
    grid_h = Inches(5.28)

    frame = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, grid_left, grid_top, grid_w, grid_h)
    frame.fill.solid()
    frame.fill.fore_color.rgb = CARD_BG
    frame.fill.transparency = 0.16
    frame.line.color.rgb = WHITE
    frame.line.transparency = 0.32
    frame.line.width = Pt(1.2)

    half_w = int(grid_w / 2)
    mid_x = grid_left + half_w
    vline = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, mid_x, grid_top, mid_x, grid_top + grid_h)
    vline.line.color.rgb = LINE
    vline.line.transparency = 0.22
    vline.line.width = Pt(1.0)

    row_count = 4
    row_h = int(grid_h / (row_count + 1))
    for idx in range(1, row_count + 1):
        y = grid_top + row_h * idx
        hline = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, grid_left, y, grid_left + grid_w, y)
        hline.line.color.rgb = LINE
        hline.line.transparency = 0.22
        hline.line.width = Pt(0.8)

    cell_w = half_w - Inches(0.3)
    _add_textbox(slide, grid_left + Inches(0.15), grid_top + Inches(0.1), cell_w, Inches(0.35), "ПЛАН", 12, CARD_TITLE, bold=True, align=PP_ALIGN.CENTER)
    _add_textbox(slide, mid_x + Inches(0.15), grid_top + Inches(0.1), cell_w, Inches(0.35), "РЕАЛЬНОСТЬ", 12, CARD_TITLE, bold=True, align=PP_ALIGN.CENTER)

    rows = (_dedupe_and_shorten(bullets, max_len=92, max_items=4) + ["Цели", "Ресурсы", "Сроки", "Риски"])[:row_count]
    for idx, raw in enumerate(rows):
        plan, reality = _split_matrix_row(raw)
        y = grid_top + row_h * (idx + 1) + Inches(0.1)
        _add_textbox(slide, grid_left + Inches(0.15), y, cell_w, Inches(0.72), _ui_trim(plan, 58), 10, CARD_BODY)
        _add_textbox(slide, mid_x + Inches(0.15), y, cell_w, Inches(0.72), _ui_trim(reality, 58), 10, CARD_BODY)


def _draw_hero_concept_right(slide, right_left) -> None:
    cx = right_left + Inches(3.4)
    cy = Inches(3.8)
    for radius, trans in ((1.6, 0.72), (1.1, 0.78), (0.66, 0.82)):
        ring = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            cx - Inches(radius),
            cy - Inches(radius),
            Inches(radius * 2),
            Inches(radius * 2),
        )
        ring.fill.background()
        ring.line.color.rgb = WHITE
        ring.line.transparency = trans
        ring.line.width = Pt(2.0 if radius > 1 else 1.4)

    core = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, cx - Inches(0.22), cy - Inches(0.22), Inches(0.44), Inches(0.44))
    core.fill.solid()
    core.fill.fore_color.rgb = WHITE
    core.line.fill.background()
    tags = [
        ("Контекст", Inches(-1.55), Inches(-1.45)),
        ("Факторы", Inches(1.15), Inches(-1.2)),
        ("Решения", Inches(1.2), Inches(1.15)),
    ]
    for label, dx, dy in tags:
        bx = cx + dx
        by = cy + dy
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cx, cy, bx + Inches(0.78), by + Inches(0.22))
        conn.line.color.rgb = WHITE
        conn.line.transparency = 0.45
        conn.line.width = Pt(1.0)
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, bx, by, Inches(1.55), Inches(0.48))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.fill.transparency = 0.24
        box.line.color.rgb = WHITE
        box.line.transparency = 0.32
        box.line.width = Pt(0.8)
        _add_textbox(slide, bx + Inches(0.15), by + Inches(0.12), Inches(1.2), Inches(0.24), label, 10, CARD_TITLE, bold=True, align=PP_ALIGN.CENTER)


def _draw_cover_right(slide, right_left, slide_data: dict, title: str) -> None:
    mark_left = right_left + Inches(0.4)
    mark_top = Inches(1.1)
    mark_w = Inches(5.8)
    mark_h = Inches(5.4)
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, mark_left, mark_top, mark_w, mark_h)
    panel.fill.solid()
    panel.fill.fore_color.rgb = CARD_BG
    panel.fill.transparency = 0.22
    panel.line.color.rgb = WHITE
    panel.line.transparency = 0.42
    panel.line.width = Pt(1.0)

    _add_textbox(
        slide,
        mark_left + Inches(0.45),
        mark_top + Inches(0.95),
        mark_w - Inches(0.9),
        Inches(2.0),
        _ui_trim(title, 90, ellipsis=False),
        TS_DISPLAY,
        WHITE,
        bold=True,
        font_name=FONT_DISPLAY,
        line_spacing=1.05,
    )

    subtitle = _extract_subtitle(slide_data)
    if subtitle:
        _add_textbox(
            slide,
            mark_left + Inches(0.45),
            mark_top + Inches(3.05),
            mark_w - Inches(0.9),
            Inches(1.4),
            _ui_trim(subtitle, 160),
            TS_H2,
            BODY_ON_DARK,
            line_spacing=LINE_SPACING_BODY,
            italic=True,
        )



def _draw_agenda_right(slide, right_left, bullets: list[str]) -> None:
    items = (_dedupe_and_shorten(bullets, max_len=90, max_items=6) + [
        "Контекст и предпосылки",
        "Ключевые находки",
        "Сравнения и сценарии",
        "Выводы и следующие шаги",
    ])[:5]
    card_left = right_left + Inches(0.25)
    card_top = Inches(1.1)
    card_w = Inches(6.0)
    card_h = Inches(5.6)
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, card_left, card_top, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.fill.transparency = 0.2
    card.line.color.rgb = WHITE
    card.line.transparency = 0.36
    card.line.width = Pt(1.0)

    _add_textbox(slide, card_left + Inches(0.4), card_top + Inches(0.35), card_w - Inches(0.8), Inches(0.35), "AGENDA", 11, ACCENT_ALT, bold=True)
    _add_hairline_rule(slide, card_left + Inches(0.4), card_top + Inches(0.78), card_left + card_w - Inches(0.4), card_top + Inches(0.78), color=WHITE, transparency=0.4)

    row_h = Inches(0.9)
    start_y = card_top + Inches(1.0)
    for idx, item in enumerate(items):
        y = start_y + idx * row_h
        number_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, card_left + Inches(0.4), y, Inches(0.56), Inches(0.56))
        number_box.fill.solid()
        number_box.fill.fore_color.rgb = ACCENT if idx % 2 == 0 else ACCENT_ALT
        number_box.line.fill.background()
        _add_textbox(slide, card_left + Inches(0.4), y + Inches(0.09), Inches(0.56), Inches(0.4), f"{idx + 1:02d}", 14, WHITE, bold=True, align=PP_ALIGN.CENTER)
        heading, body = _split_label(item)
        _add_textbox(slide, card_left + Inches(1.15), y + Inches(0.02), card_w - Inches(1.55), Inches(0.34), _ui_trim(heading, 72, ellipsis=False), 14, CARD_TITLE, bold=True)
        if body and body.lower() != heading.lower():
            _add_textbox(slide, card_left + Inches(1.15), y + Inches(0.38), card_w - Inches(1.55), Inches(0.44), _ui_trim(body, 100), 10, CARD_BODY)


def _draw_context_right(slide, right_left, slide_data: dict, bullets: list[str]) -> None:
    lead_left = right_left + Inches(0.25)
    lead_top = Inches(1.1)
    lead_w = Inches(6.0)
    lead_h = Inches(5.6)
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, lead_left, lead_top, lead_w, lead_h)
    panel.fill.solid()
    panel.fill.fore_color.rgb = CARD_BG
    panel.fill.transparency = 0.22
    panel.line.color.rgb = WHITE
    panel.line.transparency = 0.38
    panel.line.width = Pt(1.0)

    _add_textbox(slide, lead_left + Inches(0.45), lead_top + Inches(0.4), lead_w - Inches(0.9), Inches(0.32), "CONTEXT", 11, ACCENT_ALT, bold=True)

    subtitle = _extract_subtitle(slide_data) or (bullets[0] if bullets else "")
    if subtitle:
        _add_textbox(slide, lead_left + Inches(0.45), lead_top + Inches(0.85), lead_w - Inches(0.9), Inches(1.6), _ui_trim(subtitle, 220), 18, CARD_TITLE, bold=True)

    body_bullets = _dedupe_and_shorten(bullets[1:] or bullets, max_len=120, max_items=3)
    for idx, item in enumerate(body_bullets):
        y = lead_top + Inches(2.75 + idx * 0.85)
        accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, lead_left + Inches(0.45), y + Inches(0.14), Inches(0.3), Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = ACCENT if idx % 2 == 0 else ACCENT_ALT
        accent.line.fill.background()
        _add_textbox(slide, lead_left + Inches(0.9), y, lead_w - Inches(1.3), Inches(0.7), _ui_trim(item, 160), 12, CARD_BODY)

    source_hint = _extract_source_hint(slide_data)
    if source_hint:
        hint_top = lead_top + lead_h - Inches(0.7)
        _add_hairline_rule(slide, lead_left + Inches(0.45), hint_top - Inches(0.1), lead_left + lead_w - Inches(0.45), hint_top - Inches(0.1), color=WHITE, transparency=0.5)
        _add_textbox(slide, lead_left + Inches(0.45), hint_top, lead_w - Inches(0.9), Inches(0.5), f"Источник: {_ui_trim(source_hint, 200)}", 9, BODY_ON_DARK)


def _draw_key_findings_right(slide, right_left, slide_data: dict, bullets: list[str]) -> None:
    stats = _extract_stats(slide_data)
    grid_left = right_left + Inches(0.2)
    grid_top = Inches(1.05)
    grid_w = Inches(6.1)
    grid_h = Inches(5.65)

    frame = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, grid_left, grid_top, grid_w, grid_h)
    frame.fill.solid()
    frame.fill.fore_color.rgb = CARD_BG
    frame.fill.transparency = 0.22
    frame.line.color.rgb = WHITE
    frame.line.transparency = 0.36
    frame.line.width = Pt(1.0)

    _add_textbox(slide, grid_left + Inches(0.4), grid_top + Inches(0.35), grid_w - Inches(0.8), Inches(0.32), "KEY FINDINGS", 11, ACCENT_ALT, bold=True)

    stat_positions = [
        (grid_left + Inches(0.35), grid_top + Inches(0.9)),
        (grid_left + Inches(3.2), grid_top + Inches(0.9)),
        (grid_left + Inches(0.35), grid_top + Inches(2.45)),
        (grid_left + Inches(3.2), grid_top + Inches(2.45)),
    ]
    stat_items = (stats + [{"value": "", "label": ""}] * 4)[:4]
    has_any_stat = any(s.get("value") or s.get("label") for s in stats[:4])
    if has_any_stat:
        for idx, stat in enumerate(stat_items):
            left, top = stat_positions[idx]
            if not (stat.get("value") or stat.get("label")):
                continue
            tile = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, Inches(2.7), Inches(1.4))
            tile.fill.solid()
            tile.fill.fore_color.rgb = WHITE
            tile.fill.transparency = 0.18
            tile.line.color.rgb = WHITE
            tile.line.transparency = 0.46
            tile.line.width = Pt(0.8)
            _add_textbox(slide, left + Inches(0.2), top + Inches(0.12), Inches(2.4), Inches(0.5), _ui_trim(stat.get("value") or "—", 18, ellipsis=False), 26, CARD_TITLE, bold=True)
            _add_textbox(slide, left + Inches(0.2), top + Inches(0.72), Inches(2.4), Inches(0.6), _ui_trim(stat.get("label") or "", 90), 10, CARD_BODY)

    body_bullets = _dedupe_and_shorten(bullets, max_len=110, max_items=3)
    list_top = grid_top + (Inches(4.05) if has_any_stat else Inches(0.95))
    for idx, item in enumerate(body_bullets):
        y = list_top + idx * Inches(0.6)
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, grid_left + Inches(0.45), y + Inches(0.17), Inches(0.14), Inches(0.14))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT if idx % 2 == 0 else ACCENT_ALT
        dot.line.fill.background()
        _add_textbox(slide, grid_left + Inches(0.75), y, grid_w - Inches(1.1), Inches(0.5), _ui_trim(item, 140), 11, CARD_BODY)


def _draw_stat_highlight_right(slide, right_left, slide_data: dict, bullets: list[str]) -> None:
    stats = _extract_stats(slide_data)
    headline_stat = stats[0] if stats else {"value": "", "label": ""}
    panel_left = right_left + Inches(0.25)
    panel_top = Inches(1.05)
    panel_w = Inches(6.0)
    panel_h = Inches(5.65)

    hero = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, panel_left, panel_top, panel_w, panel_h)
    hero.fill.solid()
    hero.fill.fore_color.rgb = CARD_BG
    hero.fill.transparency = 0.16
    hero.line.color.rgb = WHITE
    hero.line.transparency = 0.32
    hero.line.width = Pt(1.0)

    value = headline_stat.get("value") or (bullets[0] if bullets else "—")
    label = headline_stat.get("label") or (bullets[1] if len(bullets) > 1 else "Главный показатель слайда")

    _add_textbox(slide, panel_left + Inches(0.5), panel_top + Inches(0.4), panel_w - Inches(1.0), Inches(0.35), "MAIN NUMBER", TS_SMALL + 1, ACCENT_ALT, bold=True)
    _add_textbox(
        slide,
        panel_left + Inches(0.5),
        panel_top + Inches(0.85),
        panel_w - Inches(1.0),
        Inches(2.2),
        _ui_trim(value, 18, ellipsis=False),
        72,
        CARD_TITLE,
        bold=True,
        font_name=FONT_DISPLAY,
        line_spacing=0.95,
    )
    _add_textbox(
        slide,
        panel_left + Inches(0.5),
        panel_top + Inches(3.15),
        panel_w - Inches(1.0),
        Inches(1.1),
        _ui_trim(label, 160),
        TS_H2,
        CARD_BODY,
        line_spacing=LINE_SPACING_BODY,
    )

    support = _dedupe_and_shorten(bullets[1:] if headline_stat.get("value") else bullets[2:], max_len=110, max_items=2)
    for idx, item in enumerate(support):
        y = panel_top + Inches(4.45 + idx * 0.8)
        accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, panel_left + Inches(0.5), y + Inches(0.12), Inches(0.22), Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = ACCENT if idx == 0 else ACCENT_ALT
        accent.line.fill.background()
        _add_textbox(slide, panel_left + Inches(0.85), y, panel_w - Inches(1.4), Inches(0.66), _ui_trim(item, 140), 11, CARD_BODY)


def _draw_quote_highlight_right(slide, right_left, slide_data: dict, bullets: list[str]) -> None:
    quotes = _extract_quotes(slide_data)
    quote = quotes[0] if quotes else {"text": bullets[0] if bullets else "", "author": "Источник"}
    panel_left = right_left + Inches(0.3)
    panel_top = Inches(1.2)
    panel_w = Inches(5.9)
    panel_h = Inches(5.4)

    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, panel_left, panel_top, panel_w, panel_h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.fill.transparency = 0.2
    card.line.color.rgb = WHITE
    card.line.transparency = 0.34
    card.line.width = Pt(1.0)

    _add_textbox(slide, panel_left + Inches(0.45), panel_top + Inches(0.35), Inches(1.4), Inches(0.8), "“", 96, ACCENT, bold=True, font_name=FONT_DISPLAY)

    quote_text = _ui_trim(quote.get("text") or "Нет точной цитаты в источнике.", 260)
    _add_textbox(
        slide,
        panel_left + Inches(0.5),
        panel_top + Inches(1.3),
        panel_w - Inches(1.0),
        Inches(2.8),
        quote_text,
        TS_H1,
        CARD_TITLE,
        bold=False,
        font_name=FONT_DISPLAY,
        italic=True,
        line_spacing=1.25,
    )

    author = _ui_trim(quote.get("author") or "Источник", 90)
    _add_hairline_rule(slide, panel_left + Inches(0.5), panel_top + Inches(4.3), panel_left + Inches(2.5), panel_top + Inches(4.3), color=ACCENT, transparency=0.0, width=1.2)
    _add_textbox(slide, panel_left + Inches(0.5), panel_top + Inches(4.4), panel_w - Inches(1.0), Inches(0.4), author, 12, CARD_BODY, bold=True)

    support = _dedupe_and_shorten([b for b in bullets if b and b != quote.get("text")], max_len=140, max_items=2)
    for idx, item in enumerate(support):
        y = panel_top + Inches(4.95 + idx * 0.55)
        _add_textbox(slide, panel_left + Inches(0.5), y, panel_w - Inches(1.0), Inches(0.5), f"· {_ui_trim(item, 130)}", 10, CARD_BODY)


def _draw_comparison_right(slide, right_left, bullets: list[str]) -> None:
    items = _dedupe_and_shorten(bullets, max_len=140, max_items=2) + ["Позиция A", "Позиция B"]
    left_a = right_left + Inches(0.25)
    col_top = Inches(1.1)
    col_w = Inches(2.85)
    col_h = Inches(5.55)
    left_b = right_left + Inches(3.25)

    for idx, (left, heading) in enumerate([(left_a, "A"), (left_b, "B")]):
        panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, col_top, col_w, col_h)
        panel.fill.solid()
        panel.fill.fore_color.rgb = CARD_BG
        panel.fill.transparency = 0.2
        panel.line.color.rgb = WHITE
        panel.line.transparency = 0.34
        panel.line.width = Pt(1.0)
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left + Inches(0.3), col_top + Inches(0.3), Inches(0.6), Inches(0.6))
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT if idx == 0 else ACCENT_ALT
        badge.line.fill.background()
        _add_textbox(slide, left + Inches(0.3), col_top + Inches(0.4), Inches(0.6), Inches(0.4), heading, 18, WHITE, bold=True, align=PP_ALIGN.CENTER)
        head_text, body_text = _split_label(items[idx])
        _add_textbox(slide, left + Inches(1.05), col_top + Inches(0.4), col_w - Inches(1.3), Inches(0.5), _ui_trim(head_text, 80, ellipsis=False), 14, CARD_TITLE, bold=True)
        _add_textbox(slide, left + Inches(0.3), col_top + Inches(1.35), col_w - Inches(0.6), Inches(3.8), _ui_trim(body_text or head_text, 260), 11, CARD_BODY)


def _draw_takeaways_right(slide, right_left, bullets: list[str]) -> None:
    items = (_dedupe_and_shorten(bullets, max_len=140, max_items=5) + [
        "Приоритет: зафиксировать ключевые выводы.",
        "Действие: оформить дорожную карту.",
        "Контроль: назначить ответственных и KPI.",
    ])[:5]
    panel_left = right_left + Inches(0.25)
    panel_top = Inches(1.05)
    panel_w = Inches(6.0)
    panel_h = Inches(5.6)

    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, panel_left, panel_top, panel_w, panel_h)
    panel.fill.solid()
    panel.fill.fore_color.rgb = CARD_BG
    panel.fill.transparency = 0.2
    panel.line.color.rgb = WHITE
    panel.line.transparency = 0.34
    panel.line.width = Pt(1.0)

    _add_textbox(slide, panel_left + Inches(0.4), panel_top + Inches(0.35), panel_w - Inches(0.8), Inches(0.35), "TAKEAWAYS", 11, ACCENT_ALT, bold=True)
    _add_hairline_rule(slide, panel_left + Inches(0.4), panel_top + Inches(0.78), panel_left + panel_w - Inches(0.4), panel_top + Inches(0.78), color=WHITE, transparency=0.4)

    row_h = Inches(0.9)
    start_y = panel_top + Inches(1.0)
    for idx, item in enumerate(items):
        y = start_y + idx * row_h
        number_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, panel_left + Inches(0.4), y, Inches(0.58), Inches(0.58))
        number_box.fill.solid()
        number_box.fill.fore_color.rgb = ACCENT if idx % 2 == 0 else ACCENT_ALT
        number_box.line.fill.background()
        _add_textbox(slide, panel_left + Inches(0.4), y + Inches(0.12), Inches(0.58), Inches(0.4), f"{idx + 1:02d}", 14, WHITE, bold=True, align=PP_ALIGN.CENTER)
        heading, body = _split_label(item)
        _add_textbox(slide, panel_left + Inches(1.15), y, panel_w - Inches(1.5), Inches(0.36), _ui_trim(heading, 80, ellipsis=False), 13, CARD_TITLE, bold=True)
        if body and body.lower() != heading.lower():
            _add_textbox(slide, panel_left + Inches(1.15), y + Inches(0.36), panel_w - Inches(1.5), Inches(0.5), _ui_trim(body, 130), 10, CARD_BODY)


def _render_hybrid_overlay_slide(slide, slide_data: dict, slide_number: int, image_path: str | None) -> None:
    title = _extract_title(slide_data, slide_number)
    bullets = _extract_bullets(slide_data)
    layout = _layout_type(slide_data, slide_number)
    image_prompt = str(slide_data.get("image_prompt") or slide_data.get("visual_meta_prompt") or title).strip()

    _add_background_layer(slide, image_path, seed=f"{title}|{image_prompt}|{layout}")
    panel_width = _add_glass_panel(slide, width_ratio=0.36, transparency=0.32)
    _render_left_text_overlay(slide, title, bullets, panel_width, layout)

    right_left = panel_width + Inches(0.32)
    if layout == "cover":
        _draw_cover_right(slide, right_left, slide_data, title)
    elif layout == "agenda":
        _draw_agenda_right(slide, right_left, bullets)
    elif layout == "context":
        _draw_context_right(slide, right_left, slide_data, bullets)
    elif layout == "key_findings":
        _draw_key_findings_right(slide, right_left, slide_data, bullets)
    elif layout == "stat_highlight":
        _draw_stat_highlight_right(slide, right_left, slide_data, bullets)
    elif layout == "quote_highlight":
        _draw_quote_highlight_right(slide, right_left, slide_data, bullets)
    elif layout == "comparison":
        _draw_comparison_right(slide, right_left, bullets)
    elif layout == "takeaways":
        _draw_takeaways_right(slide, right_left, bullets)
    elif layout == "data_matrix":
        _draw_data_matrix_right(slide, right_left, bullets)
    elif layout == "timeline":
        _draw_timeline_right(slide, right_left, bullets)
    elif layout == "hero_concept":
        _draw_hero_concept_right(slide, right_left)
    else:
        _draw_split_infographic_right(slide, right_left, bullets)


# ─── Editorial (NotebookLM-inspired) render path ────────────────────────────
# Cream canvas, top-left navy serif headline, thin hairline rule, and a single
# editorial content area below. No AI backgrounds, no glass panels, no dual
# title columns — just type-led composition.

EDITORIAL_MARGIN_L = Inches(0.85)
EDITORIAL_MARGIN_R = Inches(0.85)
EDITORIAL_TITLE_TOP = Inches(0.55)
EDITORIAL_TITLE_W = Inches(9.6)
EDITORIAL_TITLE_H = Inches(1.55)
EDITORIAL_RULE_Y = Inches(2.35)
EDITORIAL_CONTENT_TOP = Inches(2.70)
EDITORIAL_CONTENT_H = Inches(4.20)
EDITORIAL_CONTENT_W = SLIDE_W - EDITORIAL_MARGIN_L - EDITORIAL_MARGIN_R


_KICKER_LABELS = {
    "cover": "ОБЛОЖКА",
    "agenda": "ПОВЕСТКА",
    "context": "КОНТЕКСТ",
    "key_findings": "КЛЮЧЕВЫЕ ВЫВОДЫ",
    "stat_highlight": "ЦИФРА",
    "quote_highlight": "ЦИТАТА",
    "comparison": "СРАВНЕНИЕ",
    "takeaways": "ВЫВОДЫ",
    "hero_concept": "КОНЦЕПЦИЯ",
    "timeline": "ХРОНОЛОГИЯ",
    "data_matrix": "МАТРИЦА",
    "split_infographic": "РАЗБОР",
}


def _editorial_kicker_text(layout: str, slide_number: int, total_slides: int | None = None) -> str:
    label = _KICKER_LABELS.get(layout, "РАЗДЕЛ")
    if total_slides and total_slides > 0:
        return f"{label}  ·  {slide_number:02d} / {total_slides:02d}"
    return f"{label}  ·  {slide_number:02d}"


def _editorial_headline(slide, title: str, kicker: str | None) -> None:
    if kicker:
        _add_textbox(
            slide,
            EDITORIAL_MARGIN_L,
            Inches(0.38),
            EDITORIAL_TITLE_W,
            Inches(0.22),
            kicker,
            TS_SMALL,
            SECONDARY_INK,
            bold=True,
            font_name=FONT_BASE,
            line_spacing=1.0,
        )
    _add_textbox(
        slide,
        EDITORIAL_MARGIN_L,
        EDITORIAL_TITLE_TOP + Inches(0.12),
        EDITORIAL_TITLE_W,
        EDITORIAL_TITLE_H,
        _ui_trim(title, 110, ellipsis=False),
        38,
        HEADLINE_INK,
        bold=True,
        font_name=FONT_DISPLAY,
        line_spacing=1.08,
    )
    rule = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        EDITORIAL_MARGIN_L,
        EDITORIAL_RULE_Y,
        SLIDE_W - EDITORIAL_MARGIN_R,
        EDITORIAL_RULE_Y,
    )
    rule.line.color.rgb = HAIRLINE
    rule.line.width = Pt(0.75)


def _editorial_watermark(slide) -> None:
    _add_textbox(
        slide,
        SLIDE_W - Inches(1.7),
        SLIDE_H - Inches(0.6),
        Inches(1.4),
        Inches(0.22),
        "SMARTMONEY",
        TS_MICRO,
        WATERMARK,
        bold=True,
        align=PP_ALIGN.RIGHT,
        font_name=FONT_BASE,
        line_spacing=1.0,
    )


def _editorial_bullet_list(slide, bullets: list[str], *, left=None, top=None, width=None, height=None, numbered: bool = True) -> None:
    if not bullets:
        return
    left = left if left is not None else EDITORIAL_MARGIN_L
    top = top if top is not None else EDITORIAL_CONTENT_TOP
    width = width if width is not None else EDITORIAL_CONTENT_W
    height = height if height is not None else EDITORIAL_CONTENT_H
    items = bullets[:5]
    count = len(items)
    row_h = height / max(count, 1)
    for idx, raw in enumerate(items):
        row_top = top + row_h * idx
        marker_w = Inches(0.55)
        marker_text = f"{idx + 1:02d}" if numbered else "•"
        _add_textbox(
            slide,
            left,
            row_top + Inches(0.06),
            marker_w,
            Inches(0.4),
            marker_text,
            TS_H2,
            COPPER,
            bold=True,
            font_name=FONT_DISPLAY,
            line_spacing=1.0,
        )
        _add_textbox(
            slide,
            left + marker_w,
            row_top,
            width - marker_w,
            row_h,
            _ui_trim(raw, 200),
            TS_H2,
            BODY_INK,
            font_name=FONT_BASE,
            line_spacing=1.3,
        )
        if idx < count - 1:
            sep_y = row_top + row_h - Inches(0.04)
            sep = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                left + marker_w,
                sep_y,
                left + width,
                sep_y,
            )
            sep.line.color.rgb = HAIRLINE
            sep.line.width = Pt(0.5)
            sep.line.transparency = 0.35


def _editorial_callout_card(slide, left, top, width, height, *, label: str, body: str, label_color=COPPER) -> None:
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_SURFACE
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(0.75)
    _add_textbox(
        slide,
        left + Inches(0.25),
        top + Inches(0.18),
        width - Inches(0.5),
        Inches(0.28),
        label.upper(),
        TS_SMALL,
        label_color,
        bold=True,
        font_name=FONT_BASE,
        line_spacing=1.0,
    )
    _add_textbox(
        slide,
        left + Inches(0.25),
        top + Inches(0.5),
        width - Inches(0.5),
        height - Inches(0.65),
        _ui_trim(body, 280),
        TS_BODY + 1,
        BODY_INK,
        font_name=FONT_BASE,
        line_spacing=1.28,
    )


def _draw_cover_editorial(slide, title: str, slide_data: dict) -> None:
    # The cover shows a subtitle/lead and a meta block instead of bullets.
    subtitle = _extract_subtitle(slide_data)
    if subtitle:
        _add_textbox(
            slide,
            EDITORIAL_MARGIN_L,
            EDITORIAL_CONTENT_TOP + Inches(0.1),
            EDITORIAL_CONTENT_W * 0.7,
            Inches(2.0),
            _ui_trim(subtitle, 240),
            TS_H1,
            BODY_INK,
            font_name=FONT_DISPLAY,
            italic=True,
            line_spacing=1.35,
        )
    bullets = _extract_bullets(slide_data)
    if bullets:
        lead_top = EDITORIAL_CONTENT_TOP + Inches(2.35)
        _add_textbox(
            slide,
            EDITORIAL_MARGIN_L,
            lead_top,
            EDITORIAL_CONTENT_W,
            Inches(0.25),
            "О ЧЁМ ПРЕЗЕНТАЦИЯ",
            TS_SMALL,
            SECONDARY_INK,
            bold=True,
            font_name=FONT_BASE,
            line_spacing=1.0,
        )
        lines = [f"·  {b}" for b in bullets[:3]]
        _add_textbox(
            slide,
            EDITORIAL_MARGIN_L,
            lead_top + Inches(0.32),
            EDITORIAL_CONTENT_W,
            Inches(1.4),
            "\n".join(lines),
            TS_BODY + 1,
            BODY_INK,
            font_name=FONT_BASE,
            line_spacing=1.42,
        )


def _draw_agenda_editorial(slide, title: str, slide_data: dict, bullets: list[str]) -> None:
    _editorial_bullet_list(slide, bullets, numbered=True)


def _draw_context_editorial(slide, title: str, slide_data: dict, bullets: list[str]) -> None:
    lead = _extract_subtitle(slide_data)
    content_top = EDITORIAL_CONTENT_TOP
    content_h = EDITORIAL_CONTENT_H
    if lead:
        _add_textbox(
            slide,
            EDITORIAL_MARGIN_L,
            content_top,
            EDITORIAL_CONTENT_W * 0.95,
            Inches(1.1),
            _ui_trim(lead, 220),
            TS_H1,
            BODY_INK,
            font_name=FONT_DISPLAY,
            italic=True,
            line_spacing=1.32,
        )
        content_top = content_top + Inches(1.2)
        content_h = content_h - Inches(1.2)
    _editorial_bullet_list(
        slide,
        bullets,
        top=content_top,
        height=content_h,
        numbered=False,
    )


def _draw_key_findings_editorial(slide, title: str, slide_data: dict, bullets: list[str]) -> None:
    _editorial_bullet_list(slide, bullets, numbered=True)


def _draw_takeaways_editorial(slide, title: str, slide_data: dict, bullets: list[str]) -> None:
    if not bullets:
        return
    items = bullets[:4]
    card_w = EDITORIAL_CONTENT_W / 2 - Inches(0.2)
    card_h = EDITORIAL_CONTENT_H / 2 - Inches(0.15)
    for idx, item in enumerate(items):
        col = idx % 2
        row = idx // 2
        left = EDITORIAL_MARGIN_L + col * (card_w + Inches(0.4))
        top = EDITORIAL_CONTENT_TOP + row * (card_h + Inches(0.3))
        _editorial_callout_card(
            slide,
            left,
            top,
            card_w,
            card_h,
            label=f"ВЫВОД {idx + 1:02d}",
            body=item,
        )


def _draw_stat_highlight_editorial(slide, title: str, slide_data: dict, bullets: list[str]) -> None:
    stats = _extract_stats(slide_data) or []
    if not stats:
        # fall back to bullets if structured stats are missing
        _editorial_bullet_list(slide, bullets, numbered=False)
        return
    items = stats[:3]
    col_w = EDITORIAL_CONTENT_W / len(items)
    for idx, stat in enumerate(items):
        value = str(stat.get("value") or stat.get("figure") or stat.get("number") or "").strip()
        label = str(stat.get("label") or stat.get("caption") or stat.get("title") or "").strip()
        desc = str(stat.get("description") or stat.get("note") or stat.get("context") or "").strip()
        left = EDITORIAL_MARGIN_L + idx * col_w
        _add_textbox(
            slide,
            left + Inches(0.1),
            EDITORIAL_CONTENT_TOP + Inches(0.1),
            col_w - Inches(0.2),
            Inches(1.7),
            _ui_trim(value, 18, ellipsis=False) or "—",
            72,
            COPPER,
            bold=True,
            font_name=FONT_DISPLAY,
            line_spacing=1.0,
        )
        if label:
            _add_textbox(
                slide,
                left + Inches(0.1),
                EDITORIAL_CONTENT_TOP + Inches(1.85),
                col_w - Inches(0.2),
                Inches(0.6),
                _ui_trim(label, 80, ellipsis=False),
                TS_H2,
                HEADLINE_INK,
                bold=True,
                font_name=FONT_BASE,
                line_spacing=1.2,
            )
        if desc:
            _add_textbox(
                slide,
                left + Inches(0.1),
                EDITORIAL_CONTENT_TOP + Inches(2.55),
                col_w - Inches(0.2),
                Inches(1.4),
                _ui_trim(desc, 180),
                TS_BODY,
                BODY_INK,
                font_name=FONT_BASE,
                line_spacing=1.35,
            )


def _draw_quote_highlight_editorial(slide, title: str, slide_data: dict, bullets: list[str]) -> None:
    quotes = _extract_quotes(slide_data) or []
    if quotes:
        q = quotes[0]
        quote_text = str(q.get("text") or q.get("quote") or q.get("body") or "").strip()
        author = str(q.get("author") or q.get("speaker") or q.get("source") or "").strip()
    else:
        quote_text = bullets[0] if bullets else ""
        author = ""
    if not quote_text:
        return
    _add_textbox(
        slide,
        EDITORIAL_MARGIN_L + Inches(0.1),
        EDITORIAL_CONTENT_TOP + Inches(0.1),
        Inches(0.6),
        Inches(1.2),
        "\u201C",
        96,
        COPPER,
        bold=True,
        font_name=FONT_DISPLAY,
        line_spacing=0.8,
    )
    _add_textbox(
        slide,
        EDITORIAL_MARGIN_L + Inches(0.85),
        EDITORIAL_CONTENT_TOP + Inches(0.2),
        EDITORIAL_CONTENT_W - Inches(1.0),
        Inches(3.0),
        _ui_trim(quote_text, 320),
        TS_H1 + 4,
        HEADLINE_INK,
        font_name=FONT_DISPLAY,
        italic=True,
        line_spacing=1.32,
    )
    if author:
        _add_textbox(
            slide,
            EDITORIAL_MARGIN_L + Inches(0.85),
            EDITORIAL_CONTENT_TOP + Inches(3.3),
            EDITORIAL_CONTENT_W - Inches(1.0),
            Inches(0.4),
            f"— {author}",
            TS_H2,
            SECONDARY_INK,
            font_name=FONT_BASE,
            line_spacing=1.2,
        )


def _draw_comparison_editorial(slide, title: str, slide_data: dict, bullets: list[str]) -> None:
    pairs = [_split_matrix_row(b) for b in bullets[:5]] if bullets else []
    col_w = EDITORIAL_CONTENT_W / 2 - Inches(0.2)
    for col_idx, header in enumerate(["ДО", "ПОСЛЕ"]):
        left = EDITORIAL_MARGIN_L + col_idx * (col_w + Inches(0.4))
        _add_textbox(
            slide,
            left,
            EDITORIAL_CONTENT_TOP,
            col_w,
            Inches(0.3),
            header,
            TS_SMALL,
            COPPER if col_idx == 1 else SECONDARY_INK,
            bold=True,
            font_name=FONT_BASE,
            line_spacing=1.0,
        )
        lines = []
        for pair in pairs:
            left_text, right_text = pair
            value = right_text if col_idx == 1 else left_text
            if value:
                lines.append(f"·  {value}")
        if not lines and bullets:
            lines = [f"·  {b}" for b in bullets[:4]]
        _add_textbox(
            slide,
            left,
            EDITORIAL_CONTENT_TOP + Inches(0.5),
            col_w,
            EDITORIAL_CONTENT_H - Inches(0.5),
            "\n".join(lines),
            TS_BODY + 2,
            BODY_INK,
            font_name=FONT_BASE,
            line_spacing=1.4,
        )


def _draw_fallback_editorial(slide, title: str, slide_data: dict, bullets: list[str]) -> None:
    lead = _extract_subtitle(slide_data)
    content_top = EDITORIAL_CONTENT_TOP
    content_h = EDITORIAL_CONTENT_H
    if lead:
        _add_textbox(
            slide,
            EDITORIAL_MARGIN_L,
            content_top,
            EDITORIAL_CONTENT_W * 0.95,
            Inches(1.1),
            _ui_trim(lead, 220),
            TS_H1,
            BODY_INK,
            font_name=FONT_DISPLAY,
            italic=True,
            line_spacing=1.32,
        )
        content_top = content_top + Inches(1.2)
        content_h = content_h - Inches(1.2)
    _editorial_bullet_list(slide, bullets, top=content_top, height=content_h, numbered=False)


_EDITORIAL_LAYOUT_DISPATCH = {
    "cover": _draw_cover_editorial,
    "agenda": _draw_agenda_editorial,
    "context": _draw_context_editorial,
    "key_findings": _draw_key_findings_editorial,
    "stat_highlight": _draw_stat_highlight_editorial,
    "quote_highlight": _draw_quote_highlight_editorial,
    "comparison": _draw_comparison_editorial,
    "takeaways": _draw_takeaways_editorial,
}


def _render_editorial_slide(slide, slide_data: dict, slide_number: int, total_slides: int | None) -> None:
    _set_slide_background(slide)
    title = _extract_title(slide_data, slide_number)
    bullets = _extract_bullets(slide_data)
    layout = _layout_type(slide_data, slide_number)
    kicker = _editorial_kicker_text(layout, slide_number, total_slides)

    if layout == "cover":
        # cover uses the kicker as a short tagline ("ОБЛОЖКА · 01 / 08")
        _editorial_headline(slide, title, kicker)
    else:
        _editorial_headline(slide, title, kicker)

    drawer = _EDITORIAL_LAYOUT_DISPATCH.get(layout)
    if drawer is None:
        _draw_fallback_editorial(slide, title, slide_data, bullets)
    else:
        try:
            if drawer is _draw_cover_editorial:
                drawer(slide, title, slide_data)
            else:
                drawer(slide, title, slide_data, bullets)
        except Exception as exc:
            print(f"⚠️ editorial drawer {drawer.__name__} failed: {exc}")
            _draw_fallback_editorial(slide, title, slide_data, bullets)

    _editorial_watermark(slide)


def _render_slide(slide, slide_data: dict, slide_number: int, image_path: str | None, total_slides: int | None = None) -> None:
    # image_path is ignored in the editorial redesign — kept for signature stability
    _render_editorial_slide(slide, slide_data, slide_number, total_slides)
    title = _extract_title(slide_data, slide_number)
    bullets = _extract_bullets(slide_data)
    layout = _layout_type(slide_data, slide_number)
    notes_text = _compose_notes_text(slide_data, title, layout, bullets)
    _attach_speaker_notes(slide, notes_text)


def _image_is_acceptable(path: str) -> bool:
    """Reject generated images that would hurt the slide.

    Drops files that are too small, too uniform (monotone blobs),
    or overly dark/washed-out where overlay text becomes illegible.
    """
    try:
        if not path or not os.path.exists(path):
            return False
        if os.path.getsize(path) < 10_000:
            return False
        if not _PIL_AVAILABLE:
            return True  # no Pillow → trust the source
        with Image.open(path) as img:
            if img.width < 800 or img.height < 450:
                return False
            sample = img.convert("RGB").resize((120, 68))
            stat = ImageStat.Stat(sample)
            mean = sum(stat.mean) / 3.0
            stddev = sum(stat.stddev) / 3.0
            if stddev < 8.0:  # flat/monotone
                return False
            if mean < 18 or mean > 244:  # pure black/white
                return False
        return True
    except Exception as exc:
        print(f"image quality check failed: {exc}")
        return False


SOURCE_KIND_LABELS = {
    "pdf": "PDF",
    "docx": "DOCX",
    "url": "Web",
    "youtube": "YouTube",
    "image": "Изображение",
    "video": "Видео",
    "voice": "Voice-сообщение",
    "audio": "Аудиозапись",
    "text": "Текст",
}


def _draw_source_badge(slide, meta: dict | None, slide_number: int, total_slides: int) -> None:
    source_kind = (meta or {}).get("source_kind") or "text"
    source_label = (meta or {}).get("source_label") or ""
    kind_text = SOURCE_KIND_LABELS.get(source_kind, "Текст")

    parts = [f"ПО МАТЕРИАЛУ · {kind_text.upper()}"]
    if source_label:
        trimmed = _ui_trim(str(source_label), 48, ellipsis=True)
        if trimmed:
            parts.append(trimmed)
    badge_text = "  ·  ".join(parts)

    _add_textbox(
        slide,
        Inches(0.85),
        SLIDE_H - Inches(0.42),
        Inches(9.5),
        Inches(0.24),
        badge_text,
        TS_MICRO,
        SECONDARY_INK,
        bold=True,
        font_name=FONT_BASE,
        line_spacing=1.0,
    )

    page_text = f"{slide_number:02d} / {total_slides:02d}"
    _add_textbox(
        slide,
        SLIDE_W - Inches(3.4),
        SLIDE_H - Inches(0.42),
        Inches(1.6),
        Inches(0.24),
        page_text,
        TS_MICRO,
        SECONDARY_INK,
        bold=True,
        align=PP_ALIGN.RIGHT,
        font_name=FONT_BASE,
        line_spacing=1.0,
    )


def _add_fallback_slide(prs: Presentation, slide_number: int, title: str = "Ошибка генерации") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide)
    _add_panel(slide, int(prs.slide_width * 0.12), int(prs.slide_height * 0.2), int(prs.slide_width * 0.76), int(prs.slide_height * 0.32))
    _add_textbox(slide, Inches(2.4), Inches(2.6), Inches(8.5), Inches(0.5), f"{title} · Слайд {slide_number}", 20, RGBColor(180, 30, 30), bold=True, align=PP_ALIGN.CENTER)


async def build_pptx_from_json(json_string: str, output_filename: str, meta: dict | None = None) -> str:
    generated_image_paths: list[str] = []
    try:
        cleaned_json = _clean_json_string(json_string)
        if not cleaned_json:
            raise ValueError("Empty slide JSON")

        payload = json.loads(cleaned_json)
        slides = _normalize_slides(payload)
        if not slides:
            raise ValueError("No slides found in JSON")

        output_dir = os.path.dirname(output_filename)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)

        prs = Presentation()
        prs.slide_width = SLIDE_W
        prs.slide_height = SLIDE_H
        image_success_count = 0
        rendered_count = 0

        for index, slide_data in enumerate(slides):
            slide_number = int(slide_data.get("slide_number") or index + 1)
            image_prompt = str(
                slide_data.get("image_prompt")
                or slide_data.get("visual_meta_prompt")
                or slide_data.get("title")
                or ""
            ).strip()
            # Editorial redesign: no AI-generated backgrounds.
            image_path = None

            slide = prs.slides.add_slide(prs.slide_layouts[6])
            try:
                _render_slide(slide, slide_data, slide_number, image_path, total_slides=len(slides))
                _draw_source_badge(slide, meta, slide_number, len(slides))
                rendered_count += 1
            except Exception as render_exc:
                print(f"⚠️ Slide {slide_number}: render error: {render_exc}")
                _set_slide_background(slide)
                _add_panel(
                    slide,
                    int(prs.slide_width * 0.08),
                    int(prs.slide_height * 0.16),
                    int(prs.slide_width * 0.84),
                    int(prs.slide_height * 0.68),
                    fill_color=WHITE,
                    transparency=0.04,
                )
                _add_textbox(
                    slide,
                    Inches(1.18),
                    Inches(1.12),
                    Inches(10.8),
                    Inches(0.72),
                    _extract_title(slide_data, slide_number),
                    30,
                    INK,
                    bold=True,
                )
                for idx, bullet in enumerate(_extract_bullets(slide_data)[:4]):
                    _add_textbox(slide, Inches(1.25), Inches(2.1 + idx * 0.8), Inches(10.3), Inches(0.52), f"• {bullet}", 18, MUTED)

            if index < len(slides) - 1:
                await asyncio.sleep(1)

        if len(prs.slides) == 0 or rendered_count == 0:
            raise ValueError("No slides were created")

        prs.save(output_filename)
        print(
            f"Presentation saved: rendered={rendered_count}/{len(slides)} "
            f"image_bg={image_success_count}/{len(slides)}"
        )
        return output_filename
    except Exception as exc:
        raise ValueError(f"Failed to build presentation: {exc}") from exc
    finally:
        for p in generated_image_paths:
            if os.path.exists(p):
                try:
                    os.remove(p)
                except OSError:
                    pass
